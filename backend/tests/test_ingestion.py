import copy
import io
import uuid
import zipfile
from datetime import timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.core.config import ROOT, settings
from app.db.models import DatasetVersion, Job, OrderLine, Upload, Workspace, now
from app.ingestion.normalize import allocate, normalize, timestamp
from app.ingestion.parser import ParseError, parse
from app.ingestion.schema import ImportConfig
from app.ingestion.service import commit_result, schedule
from app.jobs.queue import claim
from app.jobs.worker import run_one

SAMPLES = ROOT / "samples"


def config(kind="orders", **options):
    columns = {
        k: k
        for k in (
            ["sku", "unit_cost", "currency", "effective_from", "effective_to"]
            if kind == "costs"
            else [
                "order_id",
                "order_line_id",
                "sku",
                "product_name",
                "paid_at",
                "quantity",
                "restocked_quantity",
            ]
        )
    }
    if kind == "orders":
        columns.update(line_paid_minor="line_paid", line_refund_minor="line_refund")
    value = ImportConfig(columns=columns).model_dump(mode="json")
    value["options"].update(
        semantics_confirmed=True,
        coverage_confirmed=True,
        payment_state_mode="confirmed_all_paid",
        refund_mode="provided",
        coverage_start="2026-09-21T00:00:00+08:00",
        coverage_end_exclusive="2026-10-05T00:00:00+08:00",
        source_as_of="2026-10-07T00:00:00+08:00",
        **options,
    )
    return value


def parsed_orders():
    return parse(SAMPLES / "orders.csv", {})


def costs():
    return normalize(
        parse(SAMPLES / "costs.csv", {}), "costs", config("costs"), {}, "costhash", []
    )["rows"]


def normalize_orders(parsed=None, cfg=None, resolutions=None, cost_rows=None):
    return normalize(
        parsed or parsed_orders(),
        "orders",
        cfg or config(),
        resolutions or {},
        "samehash",
        costs() if cost_rows is None else cost_rows,
    )


def errors(result):
    return [i for i in result["issues"] if i["severity"] == "error"]


def test_csv_xlsx_exact_equivalence():
    csv_result = normalize_orders()
    xlsx_result = normalize_orders(parse(SAMPLES / "orders.xlsx", {}))
    assert not errors(csv_result)
    assert csv_result["rows"] == xlsx_result["rows"]
    assert csv_result["summary"]["cost_matched_rows"] == 6
    assert sum(r["line_paid_minor"] for r in csv_result["rows"]) == 115000
    assert sum(r["line_refund_minor"] for r in csv_result["rows"]) == 10000


def test_identifiers_and_same_sku_are_preserved_D02_D04():
    data = parsed_orders()
    data["rows"][0]["values"]["order_id"] = "0000123"
    data["rows"][1]["values"]["order_id"] = "12345678901234567890"
    result = normalize_orders(data)
    assert not errors(result)
    assert [r["order_id"] for r in result["rows"][:2]] == ["0000123", "12345678901234567890"]
    data["rows"][3]["values"]["sku"] = "A"
    assert len(normalize_orders(data)["rows"]) == 6


def test_duplicate_conflict_and_confirmation_D01_D03():
    data = parsed_orders()
    duplicate = copy.deepcopy(data["rows"][0])
    duplicate["row_number"] = 8
    data["rows"].append(duplicate)
    result = normalize_orders(data)
    assert not errors(result)
    assert result["summary"]["unacknowledged"] == 1
    result = normalize_orders(data, resolutions={"acknowledged_codes": ["DUPLICATE_ROW"]})
    assert result["summary"]["unacknowledged"] == 0
    assert len({r["line_id"] for r in result["rows"]}) == 7
    duplicate["values"]["line_paid"] = "301"
    assert any(i["code"] == "CONFLICTING_LINE_ID" for i in errors(normalize_orders(data)))


@pytest.mark.parametrize(
    "total,weights,expected", [(10000, [30, 70], [3000, 7000]), (100, [1, 1, 1], [34, 33, 33])]
)
def test_largest_remainder_D05_D06(total, weights, expected):
    result = allocate(total, [Decimal(v) for v in weights], [str(i) for i in range(len(weights))])
    assert result == expected and sum(result) == total


def test_order_scope_does_not_double_count():
    data = parsed_orders()
    data["rows"] = data["rows"][2:4]
    for r, weight in zip(data["rows"], [30, 70], strict=True):
        r["values"]["line_paid"] = "100"
        r["values"]["weight"] = str(weight)
    data["columns"].append("weight")
    cfg = config()
    cfg["options"]["amount_scope"] = "order_goods"
    cfg["columns"]["allocation_weight"] = "weight"
    result = normalize_orders(data, cfg)
    assert not errors(result)
    assert [r["line_paid_minor"] for r in result["rows"]] == [3000, 7000]


def test_zero_weight_refund_and_missing_values_D07_D08_D15():
    with pytest.raises(ValueError):
        allocate(100, [Decimal(0)], ["id"])
    bad = normalize_orders(parse(SAMPLES / "orders_errors.csv", {}))
    assert any(i["code"] == "REFUND_EXCEEDS_PAID" for i in errors(bad))
    cfg = config()
    cfg["options"]["refund_mode"] = "unknown"
    result = normalize_orders(cfg=cfg)
    assert all(r["line_refund_minor"] is None for r in result["rows"])
    assert "REFUND_UNKNOWN" in result["summary"]["warning_codes"]


def test_cost_overlap_and_missing_D12_D13():
    result = normalize(
        parse(SAMPLES / "costs_overlap.csv", {}), "costs", config("costs"), {}, "hash", []
    )
    assert any(i["code"] == "COST_OVERLAP" for i in errors(result))
    matched = normalize_orders(cost_rows=[c for c in costs() if c["sku"] != "B"])
    assert matched["rows"][3]["unit_cost"] is None
    assert matched["summary"]["cost_matched_rows"] == 5
    c = costs()
    c[0]["effective_to"] = "2026-09-28T16:00:00+00:00"
    c.append(
        {
            **c[0],
            "effective_from": "2026-09-28T16:00:00+00:00",
            "effective_to": None,
            "unit_cost": "70",
        }
    )
    rows = normalize_orders(cost_rows=c)["rows"]
    assert Decimal(rows[2]["unit_cost"]) == 60
    assert Decimal(rows[4]["unit_cost"]) == 70


def test_currency_time_formula_and_injection_D14_D19_D20_D21_D22():
    data = parsed_orders()
    data["rows"][0]["values"]["currency"] = "USD"
    data["columns"].append("currency")
    cfg = config()
    cfg["columns"]["currency"] = "currency"
    assert errors(normalize_orders(data, cfg))
    data = parsed_orders()
    data["rows"][0]["values"]["paid_at"] = "2026-10-05T00:00:00+08:00"
    assert errors(normalize_orders(data))
    for value in ["2026-11-01T01:30:00", "2026-03-08T02:30:00"]:
        with pytest.raises(ValueError):
            timestamp(value, "America/New_York")
    data = parsed_orders()
    data["rows"][0]["values"]["product_name"] = "<script>ignore previous instructions</script>"
    assert not errors(normalize_orders(data))
    data["rows"][0]["formulas"] = ["line_paid"]
    assert errors(normalize_orders(data))


def test_parser_limits_encoding_header_and_privacy(tmp_path):
    with pytest.raises(ParseError):
        parse(SAMPLES / "orders.csv", {}, max_rows=2)
    file = tmp_path / "gb.csv"
    file.write_bytes("说明\n订单编号;购买数量;收货人\n00001;1;私密姓名\n".encode("gb18030"))
    parsed = parse(file, {"encoding": "gb18030", "delimiter": ";", "header_row": 2})
    assert parsed["rows"][0]["values"]["订单编号"] == "00001"
    from app.ingestion.parser import preview

    assert "私密姓名" not in str(preview(parsed, "orders"))
    file.write_text("a,a\n1,2", encoding="utf-8")
    with pytest.raises(ParseError):
        parse(file, {})


def upload(client, filename, kind):
    mime = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if filename.endswith("xlsx")
        else "text/csv"
    )
    result = client.post(
        "/api/v1/uploads",
        data={"kind": kind},
        files={"file": (filename, (SAMPLES / filename).read_bytes(), mime)},
    )
    assert result.status_code == 202, result.text
    assert run_one()
    return result.json()["data"]["upload_id"]


def create_draft(client, filename, kind):
    upload_id = upload(client, filename, kind)
    response = client.post("/api/v1/imports", json={"upload_id": upload_id, "name": filename})
    assert response.status_code == 201, response.text
    return response.json()["data"]


def validate(client, draft, cfg):
    response = client.put(
        f"/api/v1/imports/{draft['id']}/mapping",
        json={"expected_revision": draft["revision"], "config": cfg},
    )
    assert response.status_code == 200, response.text
    draft = response.json()["data"]
    response = client.post(
        f"/api/v1/imports/{draft['id']}/validate", json={"expected_revision": draft["revision"]}
    )
    assert response.status_code == 202
    assert run_one()
    return client.get(f"/api/v1/imports/{draft['id']}").json()["data"]


def publish(client, draft):
    response = client.post(
        f"/api/v1/imports/{draft['id']}/publish", json={"expected_revision": draft["revision"]}
    )
    assert response.status_code == 202, response.text
    assert run_one()
    item = client.get(f"/api/v1/imports/{draft['id']}").json()["data"]
    assert item["status"] == "published", item
    return item["published_id"]


def test_cost_orders_publish_immutable_and_reuse(api_client, db, isolated_storage):
    cost = validate(api_client, create_draft(api_client, "costs.csv", "costs"), config("costs"))
    assert cost["status"] == "ready", cost
    cost_id = publish(api_client, cost)
    cfg = config()
    cfg["cost_version_id"] = cost_id
    draft = validate(api_client, create_draft(api_client, "orders.csv", "orders"), cfg)
    assert draft["status"] == "ready", draft
    version = publish(api_client, draft)
    snapshot = api_client.get(f"/api/v1/datasets/{version}").json()["data"]
    assert snapshot["row_count"] == 6 and snapshot["summary"]["cost_matched_rows"] == 6
    assert (
        api_client.put(
            f"/api/v1/imports/{draft['id']}/mapping",
            json={"expected_revision": draft["revision"], "config": cfg},
        ).status_code
        == 409
    )
    repeated = validate(api_client, create_draft(api_client, "orders.csv", "orders"), cfg)
    assert publish(api_client, repeated) == version
    assert db.scalar(select(func.count()).select_from(DatasetVersion)) == 1
    assert db.scalar(select(func.count()).select_from(OrderLine)) == 6
    cfg["options"]["source_as_of"] = "2026-10-07T01:00:00+08:00"
    revised = validate(api_client, create_draft(api_client, "orders.csv", "orders"), cfg)
    assert publish(api_client, revised) != version


def test_errors_exclusion_and_warning_decisions(api_client, isolated_storage):
    draft = validate(api_client, create_draft(api_client, "orders_errors.csv", "orders"), config())
    assert draft["status"] == "review_required" and draft["summary"]["errors"] == 1
    path = f"/api/v1/imports/{draft['id']}"
    assert (
        api_client.post(
            path + "/publish", json={"expected_revision": draft["revision"]}
        ).status_code
        == 409
    )
    response = api_client.put(
        path + "/resolutions",
        json={
            "expected_revision": draft["revision"],
            "excluded_rows": [8],
            "acknowledged_codes": ["COST_MISSING"],
            "reason": "错误样例已核对；成本待补齐",
        },
    )
    assert response.status_code == 200
    assert (
        api_client.post(
            path + "/publish", json={"expected_revision": response.json()["data"]["revision"]}
        ).status_code
        == 409
    )
    api_client.post(
        path + "/validate", json={"expected_revision": response.json()["data"]["revision"]}
    )
    run_one()
    draft = api_client.get(path).json()["data"]
    assert draft["status"] == "ready", draft
    assert draft["summary"]["excluded_rows"] == 1
    assert publish(api_client, draft)


def test_stale_revision_cannot_replace_current_A13(api_client, db, isolated_storage):
    draft = create_draft(api_client, "orders.csv", "orders")
    path = f"/api/v1/imports/{draft['id']}"
    api_client.post(path + "/validate", json={"expected_revision": draft["revision"]})
    old_job = claim()
    api_client.put(
        path + "/mapping", json={"expected_revision": draft["revision"], "config": config()}
    )
    assert commit_result(old_job, {"error": "stale failure"})
    current = api_client.get(path).json()["data"]
    assert current["revision"] == 2 and current["status"] == "mapping_required"
    assert current["error_message"] is None


def test_upload_validation_and_formula_workbook(api_client, isolated_storage):
    assert (
        api_client.post(
            "/api/v1/uploads",
            data={"kind": "orders"},
            files={"file": ("bad.xlsx", b"not zip", "application/octet-stream")},
        ).status_code
        == 422
    )
    assert api_client.get(f"/api/v1/datasets/{uuid.uuid4()}").status_code == 404
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["order_id", "value"])
    sheet.append(["0001", "=1+1"])
    output = io.BytesIO()
    workbook.save(output)
    result = api_client.post(
        "/api/v1/uploads",
        data={"kind": "orders"},
        files={"file": ("formula.xlsx", output.getvalue(), "application/octet-stream")},
    )
    assert result.status_code == 202
    run_one()
    assert (
        api_client.get("/api/v1/uploads/" + result.json()["data"]["upload_id"]).json()["data"][
            "status"
        ]
        == "ready"
    )


def test_upload_size_guards_including_chunked(api_client, isolated_storage, monkeypatch):
    monkeypatch.setattr(settings(), "max_upload_mib", 1)
    response = api_client.post(
        "/api/v1/uploads", content=b"tiny", headers={"Content-Length": str(3 * 1024 * 1024)}
    )
    assert response.status_code == 413
    response = api_client.post(
        "/api/v1/uploads",
        data={"kind": "orders"},
        files={"file": ("large.csv", b"a" * (1024 * 1024 + 1), "text/csv")},
    )
    assert response.status_code == 413
    assert not list(settings().path(settings().storage_root).rglob("source.csv"))
    response = api_client.post(
        "/api/v1/uploads",
        content=iter([b"a" * (2 * 1024 * 1024 + 1)]),
        headers={"Content-Type": "multipart/form-data; boundary=test"},
    )
    assert response.status_code == 413


def test_same_revision_validation_and_late_result_cannot_unlock_publish(
    api_client, isolated_storage
):
    draft = create_draft(api_client, "orders.csv", "orders")
    path = f"/api/v1/imports/{draft['id']}"
    body = {"expected_revision": draft["revision"]}
    assert api_client.post(path + "/validate", json=body).status_code == 202
    assert api_client.post(path + "/validate", json=body).status_code == 409
    job = claim()
    # A finished job cannot submit again, even before the import revision changes.
    assert commit_result(job, {"error": "test failure"})
    assert not commit_result(job, {"error": "late result"})
    assert api_client.get(path).json()["data"]["error_message"] == "test failure"


def test_late_owned_validation_cannot_change_published_import(api_client, db, isolated_storage):
    draft = validate(api_client, create_draft(api_client, "costs.csv", "costs"), config("costs"))
    upload_row = db.get(Upload, uuid.UUID(draft["upload_id"]))
    # Simulate a duplicate from an older worker deployment; normal API rejects duplicates.
    schedule(
        db,
        upload_row.workspace_id,
        "validate_import",
        {"import_id": draft["id"], "revision": draft["revision"]},
    )
    db.commit()
    late_job = claim()
    version_id = publish(api_client, draft)
    assert commit_result(late_job, {"error": "late owned failure"})
    current = api_client.get(f"/api/v1/imports/{draft['id']}").json()["data"]
    assert current["status"] == "published" and current["published_id"] == version_id
    assert current["error_message"] is None
    assert (
        api_client.get(f"/api/v1/cost-versions/{version_id}").json()["data"]["preview"][0][
            "unit_cost"
        ]
        == "60.000000"
    )


def test_worker_failure_and_exhausted_lease_surface_to_user(
    api_client, db, isolated_storage, monkeypatch
):
    from app.ingestion import service

    draft = create_draft(api_client, "orders.csv", "orders")
    path = f"/api/v1/imports/{draft['id']}"
    response = api_client.post(path + "/validate", json={"expected_revision": draft["revision"]})
    monkeypatch.setattr(
        service, "execute", lambda job: (_ for _ in ()).throw(RuntimeError("private"))
    )
    assert run_one()
    data = api_client.get(path).json()["data"]
    assert data["status"] == "failed" and "private" not in data["error_message"]
    response = api_client.post(path + "/validate", json={"expected_revision": draft["revision"]})
    job = db.get(Job, uuid.UUID(response.json()["data"]["job_id"]))
    job.status, job.attempt = "running", settings().job_max_attempts
    job.lease_expires_at = now() - timedelta(seconds=1)
    db.commit()
    assert claim() is None
    assert api_client.get(path).json()["data"]["status"] == "failed"


def test_imports_remain_workspace_scoped(api_client, db, isolated_storage):
    upload_id = upload(api_client, "orders.csv", "orders")
    foreign = Workspace(name="Another workspace", slug="foreign")
    db.add(foreign)
    db.flush()
    db.get(Upload, uuid.UUID(upload_id)).workspace_id = foreign.id
    db.commit()
    assert api_client.get(f"/api/v1/uploads/{upload_id}").status_code == 404
    assert (
        api_client.post(
            "/api/v1/imports", json={"upload_id": upload_id, "name": "foreign"}
        ).status_code
        == 404
    )


def test_workbook_hidden_rows_merged_cells_and_archive_budget(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["订单号", "商品实付"])
    sheet.append(["000123", 10])
    sheet.row_dimensions[2].hidden = True
    sheet.merge_cells("A3:B3")
    file = tmp_path / "hidden.xlsx"
    workbook.save(file)
    data = parse(file, {})
    assert set(data["warnings"]) == {"HIDDEN_ROWS", "MERGED_CELLS"}
    assert data["rows"][0]["values"]["订单号"] == "000123"
    with zipfile.ZipFile(file, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for i in range(1001):
            archive.writestr(str(i), "")
    with pytest.raises(ParseError, match="资源超过限制"):
        parse(file, {})


def test_status_shipping_refund_allocation_and_exact_precision():
    data = parsed_orders()
    data["rows"] = data["rows"][2:4]
    cfg = config()
    cfg["options"].update(amount_scope="order_with_shipping", refund_scope="order")
    cfg["columns"].update(allocation_weight="weight", shipping_minor="shipping")
    data["columns"] += ["weight", "shipping"]
    for row, weight in zip(data["rows"], [30, 70], strict=True):
        row["values"].update(line_paid="110", line_refund="10", weight=str(weight), shipping="10")
    result = normalize_orders(data, cfg)
    assert not errors(result)
    assert [r["line_paid_minor"] for r in result["rows"]] == [3000, 7000]
    assert [r["line_refund_minor"] for r in result["rows"]] == [300, 700]
    assert all(r["quality"]["refund_allocation"] == "estimated" for r in result["rows"])
    data = parsed_orders()
    cfg = config()
    cfg["options"].update(
        payment_state_mode="mapped", status_mapping={"paid": "paid", "pending": "unpaid"}
    )
    cfg["columns"]["payment_state"] = "state"
    data["columns"].append("state")
    for row in data["rows"]:
        row["values"]["state"] = "paid"
    data["rows"][0]["values"]["state"] = "pending"
    assert normalize_orders(data, cfg)["summary"]["unpaid_rows"] == 1
    data["rows"][1]["values"]["line_paid"] = "1.00000000000000000000000000000000001"
    assert errors(normalize_orders(data, cfg))


@pytest.mark.parametrize("resource", ["time", "memory"])
def test_isolated_parser_reaps_process_on_resource_limit(isolated_storage, monkeypatch, resource):
    import subprocess
    import sys

    from app.ingestion import service

    original = subprocess.Popen
    children = []

    def launch(*args, **kwargs):
        process = original([sys.executable, "-c", "import time; time.sleep(30)"], **kwargs)
        children.append(process)
        return process

    monkeypatch.setattr(service.subprocess, "Popen", launch)
    if resource == "time":
        monkeypatch.setattr(settings(), "parse_timeout_seconds", 0.01)
    else:
        monkeypatch.setattr(
            service.psutil,
            "Process",
            lambda pid: SimpleNamespace(
                memory_info=lambda: SimpleNamespace(rss=800 * 1024 * 1024),
                children=lambda **kw: [],
                is_running=lambda: True,
            ),
        )
    result = service.isolated({}, uuid.uuid4())
    assert "超时" in result["error"] if resource == "time" else "内存超过" in result["error"]
    assert children[0].poll() is not None
    assert not list(settings().path(settings().storage_root).rglob("*.json"))
