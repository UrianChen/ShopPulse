import csv
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from defusedxml import ElementTree
from openpyxl import load_workbook

from app.ingestion.schema import COST_FIELDS, ORDER_FIELDS


class ParseError(ValueError):
    pass


def check_xlsx(path: Path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(x.file_size for x in entries) > 200 * 1024 * 1024:
                raise ParseError("XLSX 解压资源超过限制，请拆分文件。")
            names = {item.filename for item in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise ParseError("文件不是有效的 XLSX 工作簿。")
            if any("vbaproject" in n.lower() or "externallinks/" in n.lower() for n in names):
                raise ParseError("不支持包含宏或外部链接的工作簿，请转换为普通值文件。")
            for item in entries:
                if item.flag_bits & 1 or item.file_size > 100 * 1024 * 1024:
                    raise ParseError("工作簿被加密或单个条目过大。")
    except zipfile.BadZipFile:
        raise ParseError("XLSX 文件签名或压缩结构无效。") from None


def parse(path: Path, config: dict, max_rows: int = 100000) -> dict:
    header_row = config.get("header_row", 1)
    warnings = []
    sheets = []
    workbook = None
    if path.suffix == ".xlsx":
        check_xlsx(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheets = workbook.sheetnames
        if len(sheets) > 50:
            raise ParseError("最多支持 50 个工作表。")
        sheet = config.get("sheet_name") or sheets[0]
        if sheet not in sheets:
            raise ParseError("选定的工作表不存在。")
        worksheet = workbook[sheet]
        if worksheet.sheet_state != "visible":
            warnings.append("HIDDEN_SHEET")
        # Inspect XML in streaming mode without evaluating formulas or loading external content.
        with zipfile.ZipFile(path) as archive:
            with archive.open(worksheet._worksheet_path) as source:
                for _, element in ElementTree.iterparse(source, events=("end",)):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "mergeCell":
                        warnings.append("MERGED_CELLS")
                    if tag == "row" and element.attrib.get("hidden") == "1":
                        warnings.append("HIDDEN_ROWS")
                    element.clear()
        worksheet.reset_dimensions()

        def excel_rows():
            for cells in worksheet.iter_rows():
                values, formulas, numeric = [], [], []
                for index, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type == "f":
                        formulas.append(index)
                    if isinstance(value, (float, int)) and not isinstance(value, bool):
                        numeric.append(index)
                        if re.fullmatch(r"0{2,30}", cell.number_format or ""):
                            value = str(int(value)).zfill(len(cell.number_format))
                    if isinstance(value, (datetime, date)):
                        value = value.isoformat()
                    values.append("" if value is None else str(value).strip())
                yield values, formulas, numeric

        iterator = excel_rows()
    else:
        if path.read_bytes()[:4].startswith((b"PK", b"\xd0\xcf", b"\x00")):
            raise ParseError("CSV 文件类型与扩展名不符。")
        sheet = None

        def csv_rows():
            try:
                with path.open(encoding=config.get("encoding", "utf-8-sig"), newline="") as source:
                    csv.field_size_limit(4096)
                    for row in csv.reader(
                        source, delimiter=config.get("delimiter", ","), strict=True
                    ):
                        yield [value.strip() for value in row], [], []
            except (UnicodeError, csv.Error):
                raise ParseError("CSV 编码、分隔符或单元格长度不正确，请调整解析选项。") from None

        iterator = csv_rows()
    columns, rows = [], []
    try:
        for row_number, (values, formulas, numeric) in enumerate(iterator, 1):
            if row_number > max_rows + header_row:
                raise ParseError("记录行数超过限制；文件未截断，请拆分后重试。")
            if len(values) > 100 or any(len(v) > 4096 or "\x00" in v for v in values):
                raise ParseError("超过 100 列、单元格 4096 字符限制或含有二进制内容。")
            if row_number < header_row:
                continue
            if row_number == header_row:
                columns = values
                if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
                    raise ParseError("标题行不能为空或包含重复列名，请选择正确标题行。")
                if formulas:
                    raise ParseError("标题行不能包含公式。")
                continue
            if not any(values):
                continue
            if len(values) > len(columns):
                raise ParseError("数据列数超过标题列数，请检查分隔符或标题行。")
            values += [""] * (len(columns) - len(values))
            rows.append(
                {
                    "row_number": row_number,
                    "values": dict(zip(columns, values, strict=True)),
                    "formulas": [columns[x] for x in formulas],
                    "numeric": [columns[x] for x in numeric],
                }
            )
    finally:
        if workbook:
            workbook.close()
    if not columns or not rows:
        raise ParseError("选定的表格没有可读取的数据行。")
    return {
        "columns": columns,
        "rows": rows,
        "sheets": sheets,
        "sheet_name": sheet,
        "warnings": sorted(set(warnings)),
        "row_count": len(rows),
    }


def preview(parsed: dict, kind: str):
    aliases = ORDER_FIELDS if kind == "orders" else COST_FIELDS
    suggestions = {}
    known = {name.lower() for key, names in aliases.items() for name in [key, *names]}
    for field, names in aliases.items():
        match = next(
            (c for c in parsed["columns"] if c.lower() in {v.lower() for v in [field, *names]}),
            None,
        )
        if match:
            suggestions[field] = {"column": match, "confidence": 1.0 if match == field else 0.9}
    return {
        "columns": parsed["columns"],
        "sheets": parsed["sheets"],
        "sheet_name": parsed["sheet_name"],
        "row_count": parsed["row_count"],
        "warnings": parsed["warnings"],
        "suggestions": suggestions,
        "preview": [
            {
                "row_number": r["row_number"],
                "values": {
                    k: v[:120] if k.lower() in known else "（未映射列已隐藏）"
                    for k, v in r["values"].items()
                },
            }
            for r in parsed["rows"][:10]
        ],
    }
