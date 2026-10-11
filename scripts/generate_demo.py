"""Generate synthetic fixtures; expected totals are independently specified, not computed."""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
HEADER = ["order_id", "order_line_id", "sku", "product_name", "paid_at", "quantity", "line_paid", "line_refund", "restocked_quantity"]
ROWS = [
    ["O0900", "L09001", "A", "基础款收纳盒", "2026-09-22T10:00:00+08:00", "3", "300.00", "0.00", "0"],
    ["O0901", "L09011", "A", "基础款收纳盒", "2026-09-25T11:00:00+08:00", "2", "200.00", "0.00", "0"],
    ["O1001", "L10011", "A", "基础款收纳盒", "2026-09-28T10:00:00+08:00", "2", "200.00", "0.00", "0"],
    ["O1001", "L10012", "B", "大号收纳箱", "2026-09-28T10:00:00+08:00", "1", "150.00", "0.00", "0"],
    ["O1002", "L10021", "A", "基础款收纳盒", "2026-09-30T15:00:00+08:00", "1", "100.00", "100.00", "1"],
    ["O1003", "L10031", "C", "旅行收纳袋", "2026-10-04T20:00:00+08:00", "2", "200.00", "0.00", "0"],
]
COST_HEADER = ["sku", "unit_cost", "currency", "effective_from", "effective_to"]
COSTS = [[sku, cost, "CNY", "2026-01-01T00:00:00+08:00", ""] for sku, cost in [("A", "60.000000"), ("B", "100.000000"), ("C", "50.000000")]]


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def main():
    folder = ROOT / "samples"
    folder.mkdir(exist_ok=True)
    write_csv(folder / "orders.csv", HEADER, ROWS)
    write_csv(folder / "costs.csv", COST_HEADER, COSTS)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单明细"
    sheet.append(HEADER)
    for row in ROWS:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="214F41")
        sheet.column_dimensions[cell.column_letter].width = 24 if cell.column != 5 else 32
    workbook.save(folder / "orders.xlsx")
    bad = [r.copy() for r in ROWS]
    bad.append(["O1004", "L10041", "A", "错误样例", "2026-10-04T12:00:00+08:00", "1", "50.00", "99.00", "0"])
    write_csv(folder / "orders_errors.csv", HEADER, bad)
    write_csv(folder / "costs_overlap.csv", COST_HEADER, COSTS + [["A", "70", "CNY", "2026-09-01T00:00:00+08:00", ""]])
    expected = {"schema_version": "1.0", "currency": "CNY", "current": {"gross_sales_minor": 65000, "refund_amount_minor": 10000, "net_sales_minor": 55000, "paid_orders": 3, "paid_units": 6, "net_cogs_minor": 32000, "gross_profit_minor": 23000}, "previous": {"gross_sales_minor": 50000, "refund_amount_minor": 0, "net_sales_minor": 50000, "paid_orders": 2, "paid_units": 5, "net_cogs_minor": 30000, "gross_profit_minor": 20000}}
    (folder / "expected_metrics.json").write_text(json.dumps(expected, indent=2), encoding="utf-8")
    print("Generated synthetic CSV/XLSX fixtures and independent expected values.")


if __name__ == "__main__":
    main()
